"""Bounded synthetic CSV/XLSX intake preserving every row and invalid value."""
import csv
from io import StringIO, BytesIO
import json
from pathlib import PurePath
from uuid import uuid4
import zipfile
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from app.core.errors import DomainError
from app.core.serialization import digest
from app.db.models import ImportBatch, ImportRow
from app.db.session import scope_query
from app.schemas.canonical import Canonical
from app.services.finance import audit, create_transaction, get, enqueue

MAX_ROWS=200


def parse_file(filename,content):
    extension=PurePath(filename).suffix.lower()
    if extension not in ('.csv','.xlsx'):raise DomainError(400,'IMPORT_FORMAT','Use a CSV or XLSX synthetic import.')
    if len(content)>2*1024*1024:raise DomainError(400,'IMPORT_TOO_LARGE','Maximum import size is 2 MiB.')
    groups=[]
    try:
        if extension=='.csv':
            rows=list(csv.reader(StringIO(content.decode('utf-8-sig')),strict=True))
            groups=[('CSV',rows)]
        else:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members=archive.infolist()
                if len(members)>100 or sum(m.file_size for m in members)>20*1024*1024 or any('vbaProject' in m.filename or 'externalLinks/' in m.filename for m in members):raise DomainError(400,'UNSAFE_WORKBOOK','Workbook macros, external links or expansion limits are unsupported.')
            workbook=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                if len(workbook.sheetnames)>10:raise DomainError(400,'IMPORT_LIMIT','At most ten sheets are supported.')
                for sheet in workbook:
                    rows=[]
                    for index,row in enumerate(sheet.iter_rows(),1):
                        if index>MAX_ROWS+1:raise DomainError(400,'IMPORT_LIMIT','At most 200 data rows are supported.')
                        rows.append([cell.value for cell in row])
                    groups.append((sheet.title,rows))
            finally:workbook.close()
    except DomainError:raise
    except Exception:raise DomainError(400,'IMPORT_PARSE','The bounded import could not be decoded.') from None
    output=[]
    for sheet,rows in groups:
        if not rows or rows[0]!=['transaction_json']:raise DomainError(400,'IMPORT_HEADER','Each sheet must have exactly one column named transaction_json.')
        for number,values in enumerate(rows[1:],2):
            raw={'transaction_json':values[0] if values and isinstance(values[0],(str,int,bool)) else repr(values[0]) if values else None,'extra_cells':[str(v) for v in values[1:]]}
            errors=[];payload=None
            try:
                if len(values)!=1 or not isinstance(values[0],str) or values[0].startswith('='):raise ValueError('Use a text JSON cell; formulas and numeric cells are rejected.')
                parsed=json.loads(values[0]);payload=Canonical.model_validate(parsed).model_dump(mode='json')
            except ValidationError as exc:
                errors=[{'field':'.'.join(str(v) for v in e['loc']),'code':e['type'],'message':e['msg']} for e in exc.errors(include_input=False,include_context=False)]
            except Exception:
                errors=[{'field':'transaction_json','code':'INVALID_JSON_CELL','message':'Expected a single text JSON cell matching the structured transaction schema.'}]
            output.append((sheet,number,raw,payload,errors))
            if len(output)>MAX_ROWS:raise DomainError(400,'IMPORT_LIMIT','At most 200 data rows are supported across the workbook.')
    if not output:raise DomainError(400,'IMPORT_EMPTY','Include at least one data row.')
    return extension[1:].upper(),output


def preview(session,identity,storage,filename,content,correlation):
    kind,rows=parse_file(filename,content);key,hashed=storage.put(identity,content)
    batch=ImportBatch(**identity.scope(),id=uuid4(),actor_id=identity.actor_id,filename=PurePath(filename).name[:160],format=kind,object_key=key,content_digest=hashed,row_count=len(rows))
    session.add(batch);session.flush()
    for sheet,number,raw,payload,errors in rows:
        session.add(ImportRow(**identity.scope(),batch_id=batch.id,sheet=sheet,row_number=number,raw_values=raw,row_digest=digest(raw),canonical_payload=payload,validation_errors=errors))
    session.flush();audit(session,identity,'IMPORT_PREVIEWED',batch.id,1,'Synthetic import validated without dropping invalid rows',correlation,{'rows':len(rows)})
    return detail(session,identity,batch.id)


def detail(session,identity,batch_id):
    batch=get(session,ImportBatch,identity,batch_id)
    rows=session.scalars(scope_query(select(ImportRow),ImportRow,identity).where(ImportRow.batch_id==batch_id).order_by(ImportRow.sheet,ImportRow.row_number)).all()
    return {'id':str(batch.id),'filename':batch.filename,'format':batch.format,'state':batch.state,'row_count':batch.row_count,'valid_count':sum(not r.validation_errors for r in rows),'invalid_count':sum(bool(r.validation_errors) for r in rows),'rows':[{'id':str(r.id),'sheet':r.sheet,'row_number':r.row_number,'raw_values':r.raw_values,'errors':r.validation_errors,'transaction_id':str(r.transaction_id) if r.transaction_id else None} for r in rows]}


def commit(session,identity,batch_id,correlation):
    batch=get(session,ImportBatch,identity,batch_id)
    if batch.state=='COMMITTED':return detail(session,identity,batch_id)
    rows=session.scalars(scope_query(select(ImportRow),ImportRow,identity).where(ImportRow.batch_id==batch_id).order_by(ImportRow.row_number)).all()
    for row in rows:
        if row.validation_errors:continue
        created=create_transaction(session,identity,row.canonical_payload,'Imported synthetic structured row',correlation)
        row.transaction_id=created['id'];session.flush()
        enqueue(session,identity,row.transaction_id,1,'Evaluate imported row',correlation,f'import:{row.id}')
    batch.state='COMMITTED';audit(session,identity,'IMPORT_COMMITTED',batch.id,1,'Valid rows committed; invalid rows retained',correlation)
    return detail(session,identity,batch_id)
