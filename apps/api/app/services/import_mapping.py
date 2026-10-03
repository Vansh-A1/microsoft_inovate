"""Explicit CSV/XLSX mapping with immutable raw/parsed/cell validation evidence."""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
from pathlib import PurePath
import zipfile
from uuid import UUID,uuid4
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from app.core.errors import DomainError
from app.core.serialization import digest, projection
from app.db.models import ImportBatch,ImportRow
from app.db.document_models import ImportCell,Document,TransactionDocument
from app.db.session import scope_query
from app.documents.normalizer import normalized_value,NormalizationError
from app.schemas.canonical import Canonical,Line,ExpenseItem
from app.services.finance import audit,get
from app.services import imports,documents


def field_schema(path):
    pieces=path.split('.')
    if len(pieces)==1 and path in Canonical.model_fields and path not in ('lines','items'):return Canonical,path
    if len(pieces)==3 and pieces[0] in ('lines','items') and pieces[1].isdecimal() and int(pieces[1])<200:
        schema=Line if pieces[0]=='lines' else ExpenseItem
        if pieces[2] in schema.model_fields:return schema,pieces[2]
    raise DomainError(422,'MAPPING_FIELD','Map only canonical observable/reference fields, with bounded line/item indexes.')


def set_field(payload,path,value):
    parts=path.split('.')
    if len(parts)==1:payload[path]=value;return
    rows=payload.setdefault(parts[0],[]);index=int(parts[1])
    while len(rows)<=index:rows.append({})
    rows[index][parts[2]]=value


def groups(filename,content):
    extension=PurePath(filename).suffix.lower()
    if extension not in ('.csv','.xlsx') or len(content)>2*1024*1024:raise DomainError(400,'IMPORT_FORMAT','Use bounded CSV or XLSX.')
    try:
        if extension=='.csv':
            if content.startswith(b'PK'):raise ValueError()
            parsed=list(csv.reader(StringIO(content.decode('utf-8-sig')),strict=True))
            return 'CSV',[('CSV',parsed)]
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members=archive.infolist()
            if len(members)>100 or sum(m.file_size for m in members)>20*1024*1024 or any(
                any(tag in m.filename for tag in ('vbaProject','externalLinks/','embeddings/')) or m.flag_bits&1 for m in members):
                raise DomainError(400,'UNSAFE_WORKBOOK','Macros, embedded objects, external links or expansion limits are unsupported.')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
        output=[]
        try:
            if len(book.sheetnames)>10:raise DomainError(400,'IMPORT_LIMIT','Maximum ten sheets.')
            for sheet in book:
                if sheet.max_column is not None and sheet.max_column>64:raise DomainError(400,'IMPORT_LIMIT','Maximum 64 columns.')
                rows=[]
                for index,row in enumerate(sheet.iter_rows(),1):
                    if index>201:raise DomainError(400,'IMPORT_LIMIT','Maximum 200 data rows.')
                    # Preserve type and exact printed formula text; never evaluate formulas.
                    rows.append([cell.value for cell in row])
                output.append((sheet.title,rows))
        finally:book.close()
        return 'XLSX',output
    except DomainError:raise
    except Exception:raise DomainError(400,'IMPORT_PARSE','Mapped import could not be decoded safely.') from None


def preview(session,identity,storage,filename,content,mapping,defaults,date_order,correlation):
    if not mapping or len(mapping)>64 or len(set(mapping.values()))!=len(mapping):raise DomainError(422,'MAPPING_INVALID','Provide 1–64 unique canonical field mappings.')
    for field in mapping.values():field_schema(field)
    kind,source_groups=groups(filename,content);prepared=[]
    for sheet,rows in source_groups:
        if not rows:continue
        header=rows[0]
        if len(header)>64 or any(not isinstance(v,str) or not v.strip() for v in header) or len(set(header))!=len(header):
            raise DomainError(422,'IMPORT_HEADER','Use unique nonempty text column headers.')
        if not set(mapping)<=set(header):raise DomainError(422,'MAPPING_COLUMN','Every mapped column must exist in each sheet.')
        for number,values in enumerate(rows[1:],2):
            if len(prepared)>=200:raise DomainError(400,'IMPORT_LIMIT','Maximum 200 data rows across sheets.')
            raw={h:projection(v) if not isinstance(v,float) else repr(v) for h,v in zip(header,values)}
            raw['_mapping']=mapping;raw['_defaults']=defaults;raw['_date_order']=date_order
            payload=deepcopy(defaults);errors=[];cells=[];currency=payload.get('currency')
            # Currency is explicit context; normalize it before any money cells.
            ordered=sorted(mapping.items(),key=lambda pair:pair[1]!='currency')
            for column,field in ordered:
                index=header.index(column);value=values[index] if index<len(values) else None
                raw_text=None if value is None else str(value);parsed=None;validation={'state':'UNKNOWN','steps':[]}
                try:
                    if value is None or value=='':raise NormalizationError('MISSING_CELL')
                    if not isinstance(value,str):raise NormalizationError('TEXT_CELL_REQUIRED')
                    if value.lstrip().startswith(('=','+','@')):raise NormalizationError('FORMULA_CELL_REJECTED')
                    _,name=field_schema(field)
                    if name.endswith('_id') or name=='id':
                        parsed=str(UUID(value));steps=['explicit_UUID']
                    else:parsed,steps=normalized_value(field,value,currency,date_order)
                    if field=='currency':currency=parsed
                    set_field(payload,field,parsed);validation={'state':'VALID','steps':steps}
                except (NormalizationError,ValueError) as exc:
                    code=exc.code if isinstance(exc,NormalizationError) else 'UUID_INVALID'
                    errors.append({'field':field,'code':code,'message':f'{column}: {code}'})
                    validation={'state':'INVALID','code':code,'steps':[]};set_field(payload,field,None)
                cells.append({'column':column,'field_path':field,'raw_value':raw_text,'parsed_value':parsed,'validation':validation})
            # All document claims in mapped imports must bind actual safely stored documents.
            source_ids=([payload.get('source_document_id')] if payload.get('branch')=='VENDOR_INVOICE' else [i.get('source_document_id') for i in payload.get('items',[])])
            for document_id in source_ids:
                doc=session.scalar(scope_query(select(Document),Document,identity).where(Document.id==UUID(document_id))) if document_id else None
                if doc is None or doc.state not in ('READY','NEEDS_INPUT','DEPENDENCY_UNAVAILABLE') or not documents.pages(session,identity,doc.id):
                    errors.append({'field':'source_document_id','code':'ACTUAL_DOCUMENT_REQUIRED','message':'An actual safely preprocessed document UUID is required; an attachment flag or fixture fact is insufficient.'})
            try:payload=Canonical.model_validate(payload).model_dump(mode='json')
            except ValidationError as exc:
                errors += [{'field':'.'.join(str(v) for v in e['loc']),'code':e['type'],'message':e['msg']} for e in exc.errors(include_input=False,include_context=False)]
                payload=None
            prepared.append((sheet,number,raw,payload,errors,cells))
    if not prepared:raise DomainError(400,'IMPORT_EMPTY','Include data rows.')
    key,sha=storage.put(identity,content)
    batch=ImportBatch(**identity.scope(),id=uuid4(),actor_id=identity.actor_id,filename=PurePath(filename).name[:160],format=kind,
        object_key=key,content_digest=sha,row_count=len(prepared))
    session.add(batch);session.flush()
    for sheet,number,raw,payload,errors,cells in prepared:
        row=ImportRow(**identity.scope(),id=uuid4(),batch_id=batch.id,sheet=sheet,row_number=number,raw_values=raw,row_digest=digest(raw),canonical_payload=payload,validation_errors=errors)
        session.add(row);session.flush()
        for cell in cells:session.add(ImportCell(**identity.scope(),row_id=row.id,batch_id=batch.id,sheet=sheet,row_number=number,**cell))
    session.flush();audit(session,identity,'IMPORT_MAPPED',batch.id,1,'Explicit mapping and immutable cell evidence captured',correlation,{'rows':len(prepared),'mapped_columns':len(mapping)})
    return imports.detail(session,identity,batch.id)


def link_import_documents(session,identity,row,transaction_id):
    if '_mapping' not in row.raw_values:return
    p=row.canonical_payload;vendor=p['branch']=='VENDOR_INVOICE'
    ids=[p.get('source_document_id')] if vendor else [i.get('source_document_id') for i in p.get('items',[])]
    for value in dict.fromkeys(ids):
        if not value:raise DomainError(422,'ACTUAL_DOCUMENT_REQUIRED','Import attachments must be actual documents.')
        doc=get(session,Document,identity,UUID(value));count=len(documents.pages(session,identity,doc.id))
        if doc.state not in ('READY','NEEDS_INPUT','DEPENDENCY_UNAVAILABLE') or not count:raise DomainError(409,'DOCUMENT_NOT_REVIEWABLE','Imported document is no longer safe to attach.')
        session.add(TransactionDocument(**identity.scope(),transaction_id=transaction_id,transaction_version=1,
            document_id=doc.id,document_version=1,role='INVOICE' if vendor else 'RECEIPT',first_page=1,last_page=count))
    session.flush()
